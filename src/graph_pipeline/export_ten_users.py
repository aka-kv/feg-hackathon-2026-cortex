"""Source-of-truth export: 10 real, held-out, verified sessions -- their
complete real workflow (every bet opened/closed, real fixture IDs, real
timestamps), the trigger point, the prediction, and the confirmed real
outcome, all traceable back to the exact raw file and rows.

"Opened" = betslip_add_bet (fixture_id populated 69.5% of the time).
"Closed" = betslip_placed_bet (fixture_id populated 100% of the time).
Deliberately NOT "viewed" -- screen_view/fortuna_screen_view/page_view have
ZERO populated fixture_id anywhere in this dataset (verified directly), so
per-fixture browsing cannot be reconstructed, only real bet-slip actions.
This is a stronger intent signal than passive viewing, not a weaker one.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import config, cooccurrence, popularity, rank, sport_segmentation
from .build_graph import build_steps
from .clean import clean as clean_pipeline
from .evaluate import RANDOM_SEED, TEST_FRACTION, split_sessions
from .nodes import annotate

YELLOW = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
GREEN = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
BOLD = Font(bold=True)

GENUINE_HITS = [
    # Refreshed after the rank.py determinism fix (explicit score/fixture-id
    # tie-break) -- two sessions that were previously "hits" only by
    # winning a non-deterministic tie dropped out (1785590038, 1785672096),
    # confirmed directly against the corrected code, not assumed. Two new
    # ones took their place. If this list and evaluate.py's live output
    # ever disagree, evaluate.py is right -- regenerate this list from it.
    "1785577599::ead389ee47756074cf06848b059b7ced9e80569c829f29a3ee3813926f920b0c",
    "1785590016::85587aa8566465386c1dfc8334e63c792a580e00d6330d01b9f0ee713d06346f",
    "1785673493::68c5de4f6150bea0cb29a826be86f7fdcd5472eb35b5c4fe979643bd064b3502",
    "1785692893::e94585b363b971f7314ffef93c4135100b5ef2858fceb1fc6d82c9a900dc0a3f",
    "1785708852::56fd153546543a43546431026270d46e67652947be5e9c17eda3fee4d970938f",
    "1785740028::85587aa8566465386c1dfc8334e63c792a580e00d6330d01b9f0ee713d06346f",
    "1785741156::b0698d6d71fdf6ca68572b02e468a8717062cbe3eacbdac5a7137e11084bf26f",
    "1785741635::7758899e5393c54e0613016cfffa3f56adc2d271f04a311af0f156e37eec6a5d",
    "1785746148::7758899e5393c54e0613016cfffa3f56adc2d271f04a311af0f156e37eec6a5d",
    "1785756954::897a7447dfef07479b767fc985df504e2d8543716885329ccf00ac6d081e0821",
]


def build_reference_artifacts():
    """Train-only graph (proper held-out split) + full-dataset sport map,
    same construction as export_session_proof.py."""
    df, _ = clean_pipeline()
    df = annotate(df)
    steps = build_steps(df)
    all_sessions = steps["session_key"].unique().tolist()
    train_keys, test_keys = split_sessions(all_sessions, TEST_FRACTION, RANDOM_SEED)

    train_steps = steps[steps["session_key"].isin(train_keys)]
    train_df = df[df["session_key"].isin(train_keys)]
    cooc = cooccurrence.build_fixture_cooccurrence(train_steps)
    fixture_adjacency = cooccurrence.build_fixture_adjacency_index(cooc, top_k=20)
    pop = popularity.fixture_popularity_from_graph(train_steps, train_df)
    sport_map = sport_segmentation.build_fixture_sport_map(df)
    return test_keys, fixture_adjacency, pop, sport_map


def main(out_path: str):
    test_keys, fixture_adjacency, pop, sport_map = build_reference_artifacts()

    raw = pd.read_csv(config.EVENT_LOG_PATH, dtype=str, keep_default_na=False)
    raw["_ts"] = pd.to_datetime(raw["timestamp"], errors="coerce", utc=True)

    wb = Workbook()
    summary = wb.active
    summary.title = "summary"
    summary.append([
        "session_key", "real_session_id", "real_PlayerID", "in_test_split",
        "total_events", "opens_add_bet", "closes_placed_bet",
        "distinct_fixtures_opened", "trigger_fixtures_first3",
        "predicted_top3", "actual_next_fixture_hit", "verified_hit",
    ])
    for cell in summary[1]:
        cell.font = BOLD

    for sk in GENUINE_HITS:
        s, pid = sk.split("::")
        sub = raw[(raw["session"] == s) & (raw["PlayerID"] == pid)].sort_values("_ts").reset_index(drop=True)

        opens = sub[sub["event_name"] == "betslip_add_bet"]
        closes = sub[sub["event_name"] == "betslip_placed_bet"]

        # Ordered distinct fixtures first seen -- ANY fixture-bearing row,
        # not just betslip_add_bet. Matches evaluate.py's validated
        # methodology exactly: a fixture can first appear via
        # betslip_placed_bet too (e.g. add_bet rows missing fixture_id --
        # 30.5% do, see docs §3.6 rule 3). Narrowing this to add_bet-only
        # was tried first and produced a false non-hit for one of these 10
        # sessions purely from that definition mismatch, not a real miss.
        seen_order = []
        trigger_row_idx = None
        for i, row in sub.iterrows():
            fid = row["fixture_id"]
            if fid and fid not in seen_order:
                seen_order.append(fid)
                if len(seen_order) == 3 and trigger_row_idx is None:
                    trigger_row_idx = i

        prediction, hit_fixtures = [], []
        if trigger_row_idx is not None:
            result = rank.recommend(
                [f"FIXTURE:{f}" for f in seen_order[:3]], fixture_adjacency, pop,
                protection_state="allowed", top_n=3, sport_map=sport_map,
            )
            prediction = [c["fixture"].replace("FIXTURE:", "") for c in result["candidates"]]
            future = set(seen_order[3:])
            hit_fixtures = [f for f in prediction if f in future]

        summary.append([
            sk, s, pid, sk in test_keys, len(sub), len(opens), len(closes),
            len(set(seen_order)), ", ".join(seen_order[:3]),
            ", ".join(prediction), ", ".join(hit_fixtures), bool(hit_fixtures),
        ])

        # Per-user detail sheet: real raw rows, opens/closes only (the
        # columns that actually carry fixture identity), trigger row
        # highlighted yellow, real-hit rows highlighted green.
        sheet_name = f"user_{s}"[:31]
        ws = wb.create_sheet(sheet_name)
        cols = ["event_name", "timestamp", "platform", "fixture_id", "selection_id", "betslip_number", "status"]
        ws.append(cols)
        for cell in ws[1]:
            cell.font = BOLD

        action_rows = sub[sub["event_name"].isin(["betslip_add_bet", "betslip_placed_bet", "betslip_placed"])]
        for i, row in action_rows.iterrows():
            ws.append([row[c] for c in cols])
            excel_row = ws.max_row
            if i == trigger_row_idx:
                for cell in ws[excel_row]:
                    cell.fill = YELLOW
            elif row["fixture_id"] in hit_fixtures:
                for cell in ws[excel_row]:
                    cell.fill = GREEN

    wb.save(out_path)
    print(f"Saved {out_path}")

    # Also persist a lightweight JSON summary for the frontend's login
    # picker -- same data as the summary sheet, machine-readable.
    import json

    json_rows = []
    for row in summary.iter_rows(min_row=2, values_only=True):
        json_rows.append(dict(zip(
            ["session_key", "real_session_id", "real_player_id", "in_test_split",
             "total_events", "opens", "closes", "distinct_fixtures_opened",
             "trigger_fixtures", "predicted_top3", "actual_hit_fixtures", "verified_hit"],
            row,
        )))
    json_path = out_path.rsplit(".", 1)[0] + ".json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(json_rows, f, indent=2)
    print(f"Saved {json_path}")


if __name__ == "__main__":
    from pathlib import Path
    out = Path(__file__).resolve().parents[2] / "artifacts" / "ten_users_source_of_truth.xlsx"
    main(str(out))
